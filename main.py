#! /usr/bin/env python3
#  -*- coding: utf-8 -*-
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import os.path
import threading
import time
import pandas as pd
import openpyxl
import sqlite3
from openpyxl.worksheet.table import Table, TableStyleInfo

_location = os.path.dirname(__file__)

LIGHT_MODE = {
    "bg": "#f0f0f0",
    "fg": "#000000",
    "throughcolor": "#DDD",
    "background": "#EEE",
    "arrowcolor": "#333",
    "bordercolor": "#BBB",
    "bg_disabled": "#DDD",
    "bg_active": "#FFF"
}

DARK_MODE = {
    "bg": "#2e2e2e",
    "fg": "#ffffff",
    "throughcolor": "#333",
    "background": "#555",
    "arrowcolor": "#FFF",
    "bordercolor": "#444",
    "bg_disabled": "#444",
    "bg_active": "#666"
}

class MainWindow:
    def __init__(self, top=None):
        """Configures and populates the toplevel window"""

        top.geometry("600x500")
        top.title("Excel SQL Query Tool")
        top.configure(background=LIGHT_MODE["bg"])

        ## Variables

        self.root = top
        self.root.title("Excel SQL Query Tool")
        self.dialog = None

        self.top = top
        self.input_file = None
        self.output_file = None
        self.query_start_time = 0
        self.timer_running = False
        self.elapsed = 0
        self.done_loading = False
        self.cancel = False
        self.query_running = False
        self.char_list = [' ', '\n', ',', '.', '[', ']', '(', ')', '\'']
        self.current_theme = LIGHT_MODE
        self.skip_load_dialog = False

        ## GUI definition

        # Menu Bar
        self.menu_bar = tk.Menu(self.top)
        self.settings_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.settings_menu.add_command(label="Toggle Dark/Light Mode", command=self.toggle_theme)
        self.menu_bar.add_cascade(label="Settings", menu=self.settings_menu)
        self.top.configure(menu=self.menu_bar)

        # Input File Selection
        self.input_frame = tk.Frame(self.top, background=LIGHT_MODE["bg"])
        self.input_frame.pack(pady=10, fill="x", padx=10)

        self.input_button = tk.Button(self.input_frame, text="Select Input File", command=self.load_file)
        self.input_button.pack(side="left", padx=5)

        self.input_entry = tk.Entry(self.input_frame, width=40)
        self.input_entry.pack(side="left", fill="x", expand=True, padx=5)
        self.input_entry.bind("<Return>", lambda event, entry=self.input_entry: self.load_file(event))

        # Shows the number of sheets
        self.sheet_label = tk.Label(self.input_frame, text="Sheets: 0", background=LIGHT_MODE["bg"])
        self.sheet_label.pack(side="left", padx=5)

        # Sheet and column display
        self.sheet_listbox_frame = tk.Frame(self.top, background=LIGHT_MODE["bg"])
        self.sheet_listbox_frame.pack(pady=10, fill="x", padx=10)

        self.sheet_listbox = tk.Listbox(self.sheet_listbox_frame, height=5, width=45)
        self.sheet_listbox.pack(side="left", fill="x", padx=5)
        self.sheet_listbox.insert(tk.END, "Select a file to see sheets...")

        self.column_listbox = tk.Listbox(self.sheet_listbox_frame, height=5, width=600, selectmode=tk.SINGLE)
        self.column_listbox.pack(side="right", fill="x", padx=5)
        self.column_listbox.insert(tk.END, "Select a sheet to see columns...")

        # Bind the sheet listbox to update the columns list when a sheet is selected
        self.sheet_listbox.bind('<<ListboxSelect>>', self.on_sheet_select)

        # Output File Selection
        self.output_frame = tk.Frame(self.top, background=LIGHT_MODE["bg"])
        self.output_frame.pack(pady=5, fill="x", padx=10)

        self.output_button = tk.Button(self.output_frame, text="Select Output File", command=self.save_file)
        self.output_button.pack(side="left", padx=5)

        self.output_entry = tk.Entry(self.output_frame, width=40)
        self.output_entry.pack(side="left", fill="x", expand=True, padx=5)

        # SQL Query Section
        self.sql_frame = tk.Frame(self.top, background=LIGHT_MODE["bg"])
        self.sql_frame.pack(pady=5, fill="both", expand=True, padx=10)

        self.sql_button_frame = tk.Frame(self.sql_frame, background=LIGHT_MODE["bg"], height=10)
        self.sql_button_frame.pack(pady=5, fill="x", padx=5)

        self.load_query_button = tk.Button(self.sql_button_frame, text="Load SQL Query", command=self.load_sql_query)
        self.load_query_button.pack(side="left", fill="x")

        self.save_query_button = tk.Button(self.sql_button_frame, text="Save SQL Query", command=self.save_sql_query)
        self.save_query_button.pack(side="left", fill="x", padx=10)

        # SQL Query input box
        self.sql_text_frame = tk.Frame(self.sql_frame)
        self.sql_text_frame.pack(side="bottom", fill="both", expand=True)

        self.sql_text = tk.Text(self.sql_text_frame, wrap="word", height=6)
        self.sql_text.grid(row=0, column=0, sticky="nsew")
        self.scrollbar = ttk.Scrollbar(self.sql_text_frame, orient="vertical", command=self.sql_text.yview)
        self.sql_text.configure(yscrollcommand=self.scrollbar.set)
        self.scrollbar.grid(row=0, column=1, sticky="ns")

        self.sql_text_frame.grid_rowconfigure(0, weight=1)
        self.sql_text_frame.grid_columnconfigure(0, weight=1)

        # Add context menu
        self.sql_text.bind("<Button-3>", self.show_context_menu)

        self.context_menu = tk.Menu(self.top, tearoff=0)
        self.context_menu.add_command(label="Cut", command=self.cut_text)
        self.context_menu.add_command(label="Copy", command=self.copy_text)
        self.context_menu.add_command(label="Paste", command=self.paste_text)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="Save file", command=self.save_sql_query)
        self.context_menu.add_command(label="Open file", command=self.load_sql_query)

        # Bind the keyboard shortcuts to custom functions
        self.sql_text.bind('<Control-BackSpace>', self.delete_word_left)
        self.sql_text.bind('<Control-Delete>', self.delete_word_right)
        self.sql_text.bind("<Control-s>", self.save_query)
        self.sql_text.bind("<Control-o>", self.open_query)

        self.sql_execute_frame = tk.Frame(self.top, background=LIGHT_MODE["bg"])
        self.sql_execute_frame.pack(pady=5, fill="x", padx=10)

        # Query execution
        self.execute_button = tk.Button(self.sql_execute_frame, text="Execute Query", command=self.execute_query)
        self.execute_button.pack(side="left", pady=2, padx=5)

        self.cancel_button = tk.Button(self.sql_execute_frame, text="Cancel Query", command=self.cancel_query)
        self.cancel_button.pack(side="left", pady=2, padx=5)

        self.execution_time_label = tk.Label(self.sql_execute_frame, text="Time: 0s")
        self.execution_time_label.pack(side="right", pady=2, padx=5)

        # Apply initial theme
        self.apply_theme(self.top, self.current_theme)


    def apply_theme(self, widget, theme):
        """Applies selected Theme"""
        try:
            if not isinstance(widget, ttk.Scrollbar):
                widget.configure(bg=theme["bg"])
            if not isinstance(widget, (tk.Frame, tk.Tk, ttk.Scrollbar)):
                widget.configure(fg=theme["fg"])
            if isinstance(widget, (tk.Text, tk.Entry)): # set cursor color for text fields
                widget.configure(insertbackground=theme["fg"])
            if isinstance(widget, ttk.Scrollbar):
                style = ttk.Style()
                style.theme_use("alt")
                style.configure("Vertical.TScrollbar", troughcolor=theme["throughcolor"], background=theme["background"], arrowcolor=theme["arrowcolor"], bordercolor=theme["bordercolor"])
                style.map("Vertical.TScrollbar", background=[("disabled", theme["bg_disabled"]), ("active", theme["bg_active"]), ("!disabled", theme["background"])])
                widget.configure(style="Vertical.TScrollbar")
        except Exception as e:
            print(f"Couldn't set theme for {widget}, {e}")

        for child in widget.winfo_children():
            self.apply_theme(child, theme)


    def toggle_theme(self):
        """Switches between Light and Dark Theme"""
        if self.current_theme == LIGHT_MODE:
            self.apply_theme(self.top, DARK_MODE)
            self.current_theme = DARK_MODE
        else:
            self.apply_theme(self.top, LIGHT_MODE)
            self.current_theme = LIGHT_MODE


    def load_file(self, event=None):
        """Start file loading in a separate thread to prevent UI freeze"""
        if event:
            self.skip_load_dialog = True
        threading.Thread(target=self._load_file_thread, daemon=True).start()


    def _load_file_thread(self):
        """Loads the file in a separate thread to avoid freezing the UI"""
        if not self.skip_load_dialog:
            file = filedialog.askopenfilename(title="Select Input Excel File", filetypes=[("Excel Files", "*.xlsx;*.xls")])
        else:
            file = self.input_entry.get()

        self.done_loading = False
        self.skip_load_dialog = False

        if file:
            # Show "Loading sheets..." before starting the process
            self.top.after(0, lambda: self.sheet_listbox.delete(0, tk.END))
            self.top.after(0, lambda: self.sheet_listbox.insert(tk.END, "Loading sheets..."))

            self.input_file = file
            self.top.after(0, lambda: self.input_entry.delete(0, tk.END))
            self.top.after(0, lambda: self.input_entry.insert(0, self.input_file))

            if not self.output_file:
                input_dir = os.path.dirname(self.input_file)
                input_name = os.path.splitext(os.path.basename(self.input_file))[0]
                self.output_file = os.path.join(input_dir, f"{input_name}_output.xlsx")
                self.output_entry.delete(0, tk.END)
                self.output_entry.insert(0, self.output_file)

            try:
                xls = pd.ExcelFile(self.input_file)
                sheet_names = xls.sheet_names

                # Show the sheet count **immediately**
                self.top.after(0, lambda: self.sheet_label.config(text=f"Sheets: {len(sheet_names)}"))

                self.loaded_data = {sheet: pd.read_excel(xls, sheet_name=sheet) for sheet in xls.sheet_names}

                self.top.after(0, lambda: self.sheet_listbox.delete(0, tk.END))
                for sheet in xls.sheet_names:
                    self.top.after(0, lambda s=sheet: self.sheet_listbox.insert(tk.END, s))

                self.done_loading = True

            except Exception as e:
                self.top.after(0, lambda: messagebox.showerror("Error", f"Failed to load file: {e}"))
                self.top.after(0, lambda: self.sheet_listbox.delete(0, tk.END))
                self.top.after(0, lambda: self.sheet_listbox.insert(tk.END, f"Failed to load file"))


    def save_file(self):
        """Select an output Excel file path"""
        self.output_file = filedialog.asksaveasfilename(title="Select Output Excel File", defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
        if self.output_file:
            self.output_entry.delete(0, tk.END)
            self.output_entry.insert(0, self.output_file)


    def load_sql_query(self):
        """Load an SQL query from a file into the text widget"""
        query_file = filedialog.askopenfilename(title="Select SQL Query File", filetypes=[("Text Files", "*.txt")])
        if query_file:
            with open(query_file, 'r') as file:
                self.sql_text.delete(1.0, tk.END)
                self.sql_text.insert(tk.END, file.read())


    def save_sql_query(self):
        """Save the current query to a file"""
        query_text = self.sql_text.get(1.0, tk.END).strip()  # Get the text from the query box
        if not query_text:
            messagebox.showwarning("No Query", "Please write a query before saving.")
            return

        # Open a file dialog to choose the save location
        file_path = filedialog.asksaveasfilename(title="Save Query As", defaultextension=".txt", filetypes=[("Text Files", "*.txt")])
        if file_path:
            try:
                with open(file_path, 'w') as file:
                    file.write(query_text)  # Save the query to the file
                messagebox.showinfo("Saved", f"Query saved successfully to {file_path}")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to save query: {e}")


    def execute_query(self):
        """Execute the SQL query on the selected input file and save the result to the output file"""
        self.output_file = self.output_entry.get()

        if not self.input_file or not self.output_file or not self.sql_text.get("1.0", tk.END).strip():
            messagebox.showerror("Error", "Please fill in all fields.")
            return

        if not self.done_loading:
            messagebox.showerror("Error", "Please wait for data to load.")
            return

        self.execution_time_label.config(text="Running: 0s")  # Reset timer display
        self.query_start_time = time.time()  # Store start time

        self.timer_running = True
        threading.Thread(target=self._update_timer, daemon=True).start()  # Start timer in background
        threading.Thread(target=self._run_query_thread, daemon=True).start()  # Run query in background


    def cancel_query(self):
        """Cancels the Query"""
        if self.query_running:
            self.cancel = True


    def _update_timer(self):
        """Updates the timer display every second while query is running"""
        self.elapsed = 0
        while self.timer_running:
            self.top.after(0, lambda t=self.elapsed: self.execution_time_label.config(text=f"Running: {t}s"))
            time.sleep(1)
            if not self.timer_running:
                break
            self.elapsed += 1

        if self.query_running:
            self.execution_time_label.config(text=f"Done! Took: {self.elapsed}s")
            self.query_running = False
        else:
            self.execution_time_label.config(text=f"Query cancelled after {self.elapsed}s")


    def _run_query_thread(self):
        """Executes the SQL Query in a background thread to keep UI responsive"""
        self.query_running = True
        try:
            conn = sqlite3.connect(":memory:")

            # Load all sheets into SQLite
            for sheet, df in self.loaded_data.items():
                if not self.cancel:
                    df.to_sql(sheet, conn, if_exists="replace", index=False)
                else:
                    self.query_stop()
                    return

            query = self.sql_text.get("1.0", tk.END).strip()
            result_df = pd.read_sql_query(query, conn)

            if self.cancel:
                self.query_stop()
                return

            result_df.to_excel(self.output_file, index=False, sheet_name="SQLResults")

            # Open and adjust the workbook
            wb = openpyxl.load_workbook(self.output_file)
            ws = wb["SQLResults"]
            table_ref = f"A1:{chr(64 + len(result_df.columns))}{len(result_df) + 1}"
            table = Table(displayName="SQLTable", ref=table_ref)
            style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            ws.add_table(table)

            # Adjust column widths
            for column in ws.columns:
                if self.cancel:
                    self.query_stop()
                    return

                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception as e:
                        pass
                ws.column_dimensions[column_letter].width = max_length + 2

            wb.save(self.output_file)
            wb.close()

            self.timer_running = False
            self.show_success_dialog()

            if self.cancel:
                self.cancel = False
                messagebox.showinfo("Could not Cancel", "Query finished before it could be cancelled.")

        except Exception as e:
            messagebox.showerror("Error", f"An error occurred: {e}")


    def query_stop(self):
        """Function to handle stopping the Query"""
        self.cancel = False
        self.query_running = False
        self.timer_running = False
        messagebox.showinfo("Query Cancelled", "Query has been cancelled.")


    def show_success_dialog(self):
        """Shows a success dialog with a button to open the output file"""
        self.dialog = tk.Toplevel(self.root)
        self.dialog.geometry("300x100")
        self.dialog.title("Execution Complete")
        self.dialog.focus()

        success_label = tk.Label(self.dialog, text=f"Query executed successfully!\nTook {self.elapsed} seconds")
        success_label.pack(pady=0)

        open_button = tk.Button(self.dialog, text="Open Output File", width=15, command=self.open_output_file)
        open_button.pack(side="left", pady=0, padx=20)

        close_button = tk.Button(self.dialog, text="Close", width=15, command=self.dialog.destroy)
        close_button.pack(side="right", pady=0, padx=20)


    def open_output_file(self):
        """Opens the output file with the default program"""
        try:
            if os.name == 'nt':  # If it's Windows
                os.startfile(self.output_file)
            elif os.name == 'posix':  # If it's Mac/Linux
                subprocess.run(['open', self.output_file])
        except Exception as e:
            messagebox.showerror("Error", f"Could not open the file: {e}")
        self.dialog.destroy()


    def on_sheet_select(self, event):
        """Handle sheet selection and display columns in the second listbox"""
        try :
            selected_sheet = self.sheet_listbox.get(self.sheet_listbox.curselection())  # Get selected sheet name

            # Clear the columns listbox first
            self.column_listbox.delete(0, tk.END)

            # Fetch columns of the selected sheet
            if selected_sheet in self.loaded_data:
                columns = self.loaded_data[selected_sheet].columns
                for column in columns:
                    self.column_listbox.insert(tk.END, column)  # Insert each column into the column listbox
        except Exception as e:
            return


    def delete_word_left(self, event=None):
        """Deletes the word to the left of the cursor"""
        pos = self.sql_text.index(tk.INSERT)  # Get the current cursor position
        line, column = pos.split('.')  # Split line and column
        column = int(column)

        if self.sql_text.get(f"{line}.{column-1}", f"{line}.{column}") in self.char_list:
            return

        # Move the cursor left until we reach the beginning of the word
        while column > 0 and self.sql_text.get(f"{line}.{column-1}", f"{line}.{column}") not in self.char_list:
            column -= 1

        # Delete the word before the cursor
        self.sql_text.delete(f"{line}.{column}", pos)
        return 'break'  # Prevent the default action


    def delete_word_right(self, event=None):
        """Deletes the word to the right of the cursor"""
        pos = self.sql_text.index(tk.INSERT)  # Get the current cursor position
        line, column = pos.split('.')  # Split line and column
        column = int(column)

        # Move the cursor right until we reach the end of the word
        text = self.sql_text.get(f"{line}.{column}", "end-1c")
        for i, char in enumerate(text):
            if i == 0 and char in self.char_list:
                return
            if char in self.char_list:
                break
        else:
            i = len(text)

        # Delete the word after the cursor
        self.sql_text.delete(f"{line}.{column}", f"{line}.{column + i}")
        return 'break'  # Prevent the default action


    def save_query(self, event=None):
        """Opens save query dialog"""
        self.save_sql_query()
        pass


    def open_query(self, event=None):
        """Opens load query dialog"""
        self.load_sql_query()
        pass


    def show_context_menu(self, event):
        """Display the context menu at the cursor location"""
        # Post the context menu at the cursor position (x, y)
        self.context_menu.post(event.x_root, event.y_root)


    def cut_text(self):
        """Cut the selected text"""
        self.sql_text.event_generate("<<Cut>>")


    def copy_text(self):
        """Copy the selected text"""
        self.sql_text.event_generate("<<Copy>>")


    def paste_text(self):
        """Paste the text from clipboard"""
        self.sql_text.event_generate("<<Paste>>")



def start_up():
    root = tk.Tk()
    app = MainWindow(root)
    root.mainloop()


if __name__ == '__main__':
    start_up()
